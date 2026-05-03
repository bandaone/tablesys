"""
CHECKPOINT 1 – Template Parser Tests
=====================================

These tests verify that StructuralTemplateParser correctly extracts
session containers from Word and Excel timetable skeleton files.

Run with:
    cd backend
    pytest tests/test_template_parser.py -v
"""

import io
import os
import sys
import pytest

# Make sure the app package is importable from the tests directory
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.utils.template_parser import StructuralTemplateParser, SESSION_KEYWORDS


# ---------------------------------------------------------------------------
# Helpers to build in-memory test files
# ---------------------------------------------------------------------------

def _make_excel_template_bytes() -> bytes:
    """
    Build a minimal Excel template that mimics the Engineering timetable
    structure. Returns raw bytes of a .xlsx file.

    Layout (first sheet):
        Row 0: Headers  ["HOURS",  "2ND YEAR",        "",           "3RD YEAR", ""]
        Row 1: Dept      ["",       "GEN-2",          "",           "AEN",      "EEE"]
        Row 2: 07:00-08  ["07:00-08:00", "Lecture",  "",           "",         ""]
        Row 3: 08:00-09  ["08:00-09:00", "",         "",           "Lab",      ""]
        Row 4: 09:00-10  ["09:00-10:00", "",         "",           "",         "Tutorial"]
        Row 5: 14:00-17  ["14:00-17:00", "",         "Lab",        "",         ""]
    """
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active

    data = [
        ["HOURS",        "2ND YEAR",  "",      "3RD YEAR", ""],
        ["",             "GEN-2",     "",      "AEN",      "EEE"],
        ["07:00-08:00",  "Lecture",   "",      "",         ""],
        ["08:00-09:00",  "",          "",      "Lab",      ""],
        ["09:00-10:00",  "",          "",      "",         "Tutorial"],
        ["14:00-17:00",  "",          "Lab",   "",         ""],
    ]
    for row in data:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx_template_bytes() -> bytes:
    """
    Build a minimal Word document with one table that mimics a simple
    two-day timetable.

    Structure:
        MONDAY label row
        Headers: HOURS | 3RD YEAR EEE | 5TH YEAR AEN
        08:00-10:00 | Lecture          | Lab
        10:00-12:00 | Tutorial         |
        TUESDAY label row
        08:00-10:00 |                  | Lecture
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    table = doc.add_table(rows=0, cols=3)

    def add_row(cells):
        row = table.add_row()
        for i, text in enumerate(cells):
            row.cells[i].text = str(text)

    add_row(["MONDAY", "", ""])
    add_row(["HOURS",       "3RD YEAR EEE",  "5TH YEAR AEN"])
    add_row(["08:00-10:00", "Lecture",        "Lab"])
    add_row(["10:00-12:00", "Tutorial",       ""])
    add_row(["TUESDAY", "", ""])
    add_row(["08:00-10:00", "",               "Lecture"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def excel_template_path(tmp_path):
    p = tmp_path / "template.xlsx"
    p.write_bytes(_make_excel_template_bytes())
    return str(p)


@pytest.fixture
def docx_template_path(tmp_path):
    p = tmp_path / "template.docx"
    p.write_bytes(_make_docx_template_bytes())
    return str(p)


# ---------------------------------------------------------------------------
# Excel tests
# ---------------------------------------------------------------------------

class TestExcelParser:
    def test_parses_without_error(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        assert "containers" in result
        assert "shape" in result

    def test_finds_lecture_container(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        types = [c["session_type"] for c in result["containers"]]
        assert "lecture" in types, f"Expected 'lecture' in containers, got {result['containers']}"

    def test_finds_lab_container(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        types = [c["session_type"] for c in result["containers"]]
        assert "practical" in types, f"Expected 'practical' in containers, got {result['containers']}"

    def test_finds_tutorial_container(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        types = [c["session_type"] for c in result["containers"]]
        assert "tutorial" in types, f"Expected 'tutorial' in containers, got {result['containers']}"

    def test_containers_have_correct_hours(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        for c in result["containers"]:
            assert c["start_hour"] >= 7, "start_hour should be >= 7"
            assert c["end_hour"] > c["start_hour"], "end_hour must be after start_hour"
            assert c["duration"] == c["end_hour"] - c["start_hour"]

    def test_lecture_at_correct_hour(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        lectures = [c for c in result["containers"] if c["session_type"] == "lecture"]
        assert any(c["start_hour"] == 7 for c in lectures), (
            "Lecture at 07:00 not found. Containers: " + str(result["containers"])
        )

    def test_group_labels_assigned(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        for c in result["containers"]:
            assert c["group_label"], f"Container missing group_label: {c}"

    def test_shape_contains_time_col(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        assert "time_col_index" in result["shape"]

    def test_total_containers_count(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser.parse()
        # Template has sessions spread across 3 identified group columns (GEN-2, AEN, EEE)
        # Each uniquely labelled column contributes at least 1 container
        assert len(result["containers"]) >= 3, (
            f"Expected at least 3 containers, got {len(result['containers'])}: {result['containers']}"
        )


# ---------------------------------------------------------------------------
# Word (.docx) tests
# ---------------------------------------------------------------------------

class TestDocxParser:
    def test_parses_without_error(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        assert "containers" in result

    def test_finds_all_session_types(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        types = {c["session_type"] for c in result["containers"]}
        # Template has lecture, lab, tutorial
        assert "lecture" in types
        assert "practical" in types
        assert "tutorial" in types

    def test_day_assignment_monday(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        monday_containers = [c for c in result["containers"] if c["day"] == "Monday"]
        assert len(monday_containers) >= 3, (
            f"Expected >=3 containers on Monday, got {monday_containers}"
        )

    def test_day_assignment_tuesday(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        tuesday_containers = [c for c in result["containers"] if c["day"] == "Tuesday"]
        assert len(tuesday_containers) >= 1, (
            f"Expected >=1 container on Tuesday, got {tuesday_containers}"
        )

    def test_multi_hour_duration(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        # The 08:00-10:00 slots should have duration=2
        multi = [c for c in result["containers"] if c["duration"] >= 2]
        assert len(multi) > 0, "No multi-hour blocks found"

    def test_group_labels_from_headers(self, docx_template_path):
        parser = StructuralTemplateParser(docx_template_path, "docx")
        result = parser.parse()
        labels = {c["group_label"] for c in result["containers"]}
        # Should have found EEE and AEN groups from header "3RD YEAR EEE" / "5TH YEAR AEN"
        assert any("EEE" in lbl for lbl in labels), f"EEE not found in labels: {labels}"
        assert any("AEN" in lbl for lbl in labels), f"AEN not found in labels: {labels}"


# ---------------------------------------------------------------------------
# Session keyword classification tests
# ---------------------------------------------------------------------------

class TestKeywordClassification:
    """Ensure SESSION_KEYWORDS cover expected synonyms."""

    @pytest.mark.parametrize("word,expected", [
        ("Lecture",   "lecture"),
        ("LECTURE",   "lecture"),
        ("lect",      "lecture"),
        ("Lab",       "practical"),
        ("LAB",       "practical"),
        ("Practical", "practical"),
        ("prac",      "practical"),
        ("Tutorial",  "tutorial"),
        ("TUT",       "tutorial"),
    ])
    def test_session_keyword_canonical(self, word, expected, excel_template_path):
        """_classify_session must return the canonical type for each keyword."""
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser._classify_session(word)
        assert result == expected, f"'{word}' → expected '{expected}', got '{result}'"

    def test_empty_cell_returns_none(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        assert parser._classify_session("") is None

    def test_random_text_returns_none(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        assert parser._classify_session("CEE 3111 NLR2") is None


# ---------------------------------------------------------------------------
# Time range parsing tests
# ---------------------------------------------------------------------------

class TestTimeRangeParsing:
    @pytest.mark.parametrize("text,expected", [
        ("07:00-08:00",    (7, 8)),
        ("07.00 - 08.00",  (7, 8)),
        ("14:00-17:00",    (14, 17)),
        ("08:00-10:00",    (8, 10)),
        ("7:00",           (7, 8)),
    ])
    def test_parse_time_range(self, text, expected, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        result = parser._parse_time_range(text)
        assert result == expected, f"Time '{text}' → expected {expected}, got {result}"

    def test_invalid_time_returns_none(self, excel_template_path):
        parser = StructuralTemplateParser(excel_template_path, "xlsx")
        assert parser._parse_time_range("No time here") is None


# ---------------------------------------------------------------------------
# Unsupported file type
# ---------------------------------------------------------------------------

class TestValidation:
    def test_unsupported_file_type_raises(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            StructuralTemplateParser("some.pdf", "pdf")
