"""
Standalone Checkpoint 1 verifier for the StructuralTemplateParser.

Run this directly (no pytest, no conftest, no DB needed):
    cd backend
    python tests/run_template_parser_tests.py
"""

import io
import sys
import os
import traceback

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.utils.template_parser import StructuralTemplateParser

PASS = "[PASS]"
FAIL = "[FAIL]"
results = {"passed": 0, "failed": 0}


def check(name: str, condition: bool, detail: str = ""):
    if condition:
        print(f"  {PASS} {name}")
        results["passed"] += 1
    else:
        print(f"  {FAIL} {name}" + (f"\n       detail: {detail}" if detail else ""))
        results["failed"] += 1


# ---------------------------------------------------------------------------
# Build in-memory Excel template
# ---------------------------------------------------------------------------

def make_excel_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    data = [
        ["HOURS",        "2ND YEAR",  "",      "3RD YEAR", ""],
        ["",             "GEN-2",     "",      "AEN",      "EEE"],
        ["07:00-08:00",  "Lecture",   "",      "",         ""],
        ["08:00-09:00",  "",          "",      "Lab",      ""],
        ["09:00-10:00",  "",          "",      "",         "Tutorial"],
        ["14:00-17:00",  "",          "Lab",   "",         ""],
        ["TUESDAY",      "",          "",      "",         ""],
        ["08:00-09:00",  "Lecture",   "",      "Tutorial", ""],
    ]
    for row in data:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx_bytes() -> bytes:
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=0, cols=3)

    def add_row(cells):
        row = table.add_row()
        for i, text in enumerate(cells):
            row.cells[i].text = str(text)

    add_row(["MONDAY", "", ""])
    add_row(["HOURS",       "3RD YEAR EEE",  "5TH YEAR AEN"])
    add_row(["08:00-10:00", "Lecture",        "Lab"])
    add_row(["10:00-12:00", "Tutorial",       ""])
    add_row(["TUESDAY", "", ""])
    add_row(["08:00-10:00", "",               "Lecture"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------------

def containers_with(result, session_type=None, day=None):
    cs = result["containers"]
    if session_type:
        cs = [c for c in cs if c["session_type"] == session_type]
    if day:
        cs = [c for c in cs if c["day"] == day]
    return cs


# ---------------------------------------------------------------------------
# Excel tests
# ---------------------------------------------------------------------------

def test_excel(tmp_path):
    print("\n=== Excel Template Tests ===")
    p = tmp_path / "template.xlsx"
    p.write_bytes(make_excel_bytes())
    parser = StructuralTemplateParser(str(p), "xlsx")

    try:
        result = parser.parse()
    except Exception as e:
        print(f"  {FAIL} parse() raised: {e}")
        traceback.print_exc()
        results["failed"] += 5
        return

    check("parse() returns containers key",           "containers" in result)
    check("parse() returns shape key",                "shape" in result)
    check("lecture container found",                  len(containers_with(result, "lecture")) > 0,
          str(result["containers"]))
    check("lab/practical container found",            len(containers_with(result, "practical")) > 0,
          str(result["containers"]))
    check("tutorial container found",                 len(containers_with(result, "tutorial")) > 0,
          str(result["containers"]))
    check("total containers >= 5",                   len(result["containers"]) >= 5,
          f"got {len(result['containers'])}")
    check("all containers have valid start_hour",     all(c["start_hour"] >= 7 for c in result["containers"]),
          str([c["start_hour"] for c in result["containers"]]))
    check("all containers have end > start",          all(c["end_hour"] > c["start_hour"] for c in result["containers"]))
    check("duration matches end - start",             all(c["duration"] == c["end_hour"] - c["start_hour"]
                                                          for c in result["containers"]))
    check("group labels assigned",                    all(c["group_label"] for c in result["containers"]),
          str([c["group_label"] for c in result["containers"]]))
    check("time_col_index in shape",                  "time_col_index" in result["shape"])
    check("Monday lecture at 07:00",                  any(c["start_hour"] == 7 and c["session_type"] == "lecture"
                                                          for c in containers_with(result, day="Monday")),
          str(containers_with(result, day="Monday")))
    check("Tuesday containers present",               len(containers_with(result, day="Tuesday")) > 0,
          str(result["containers"]))


# ---------------------------------------------------------------------------
# Word (.docx) tests
# ---------------------------------------------------------------------------

def test_docx(tmp_path):
    print("\n=== Word (.docx) Template Tests ===")
    p = tmp_path / "template.docx"
    p.write_bytes(make_docx_bytes())
    parser = StructuralTemplateParser(str(p), "docx")

    try:
        result = parser.parse()
    except Exception as e:
        print(f"  {FAIL} parse() raised: {e}")
        traceback.print_exc()
        results["failed"] += 5
        return

    check("parse() succeeds",                         "containers" in result)
    check("lecture found",                            len(containers_with(result, "lecture")) > 0,
          str(result["containers"]))
    check("practical (lab) found",                   len(containers_with(result, "practical")) > 0,
          str(result["containers"]))
    check("tutorial found",                          len(containers_with(result, "tutorial")) > 0,
          str(result["containers"]))
    check("Monday containers >= 3",                  len(containers_with(result, day="Monday")) >= 3,
          str(containers_with(result, day="Monday")))
    check("Tuesday containers >= 1",                 len(containers_with(result, day="Tuesday")) >= 1,
          str(containers_with(result, day="Tuesday")))
    check("multi-hour blocks present",               any(c["duration"] >= 2 for c in result["containers"]),
          str([c["duration"] for c in result["containers"]]))
    labels = {c["group_label"] for c in result["containers"]}
    check("EEE group label detected",                any("EEE" in lbl for lbl in labels), str(labels))
    check("AEN group label detected",                any("AEN" in lbl for lbl in labels), str(labels))


# ---------------------------------------------------------------------------
# Session keyword and time parsing unit tests
# ---------------------------------------------------------------------------

def test_unit(tmp_path):
    print("\n=== Unit Tests (keywords, time parsing) ===")
    p = tmp_path / "dummy.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Hours", "GEN-2"])
    buf = io.BytesIO()
    wb.save(buf)
    p.write_bytes(buf.getvalue())

    parser = StructuralTemplateParser(str(p), "xlsx")

    kw_cases = [
        ("Lecture",   "lecture"),
        ("LECTURE",   "lecture"),
        ("lect",      "lecture"),
        ("Lab",       "practical"),
        ("Practical", "practical"),
        ("prac",      "practical"),
        ("Tutorial",  "tutorial"),
        ("TUT",       "tutorial"),
    ]
    for word, expected in kw_cases:
        check(f"keyword '{word}' -> '{expected}'", parser._classify_session(word) == expected)

    check("empty string -> None",   parser._classify_session("") is None)
    check("course code -> None",    parser._classify_session("CEE 3111 NLR2") is None)

    time_cases = [
        ("07:00-08:00", (7,  8)),
        ("07.00-08.00", (7,  8)),
        ("14:00-17:00", (14, 17)),
        ("08:00-10:00", (8,  10)),
        ("7:00",        (7,  8)),
    ]
    for text, expected in time_cases:
        check(f"time '{text}' -> {expected}", parser._parse_time_range(text) == expected)

    check("invalid time -> None",   parser._parse_time_range("No time") is None)


# ---------------------------------------------------------------------------
# Validation test
# ---------------------------------------------------------------------------

def test_validation():
    print("\n=== Validation Tests ===")
    try:
        StructuralTemplateParser("some.pdf", "pdf")
        check("pdf raises ValueError", False, "No exception raised")
    except ValueError:
        check("pdf raises ValueError", True)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import tempfile, pathlib
    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td)
        test_excel(tmp)
        test_docx(tmp)
        test_unit(tmp)
        test_validation()

    total = results["passed"] + results["failed"]
    bar = "=" * 50
    print(f"\n{bar}")
    print(f"CHECKPOINT 1 RESULTS: {results['passed']}/{total} passed")
    if results["failed"] > 0:
        print(f"  {results['failed']} FAILURES — review output above")
        sys.exit(1)
    else:
        print("  ALL TESTS PASSED")
        sys.exit(0)
